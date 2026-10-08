"""Bounded structured table parsing and deterministic operations; never evaluates code."""
from collections import defaultdict
from decimal import Decimal, InvalidOperation, localcontext
import io
import json
import re
import time
import zipfile
from defusedxml import ElementTree as SafeET

MAX_CELLS = 100_000
MAX_PHYSICAL_CELLS = 200_000
MAX_SHEETS = 20
MAX_COLUMNS = 100

def workbook_sections(data):
    import openpyxl
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    started = time.monotonic()
    coordinates = {}
    useful = physical = 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(x.file_size for x in archive.infolist()) > 20*1024*1024:
            raise ValueError("Workbook expands beyond 20 MiB.")
        if any("vbaproject" in n.lower() for n in archive.namelist()):
            raise ValueError("Macro-enabled workbooks are not supported.")
        sheets = [n for n in archive.namelist() if n.startswith("xl/worksheets/") and n.endswith(".xml")]
        if len(sheets) > MAX_SHEETS:
            raise ValueError("Workbook exceeds 20 sheets.")
        # Reject worksheet relationship paths that would bypass the bounded
        # preflight. External links are never followed by this parser.
        import posixpath
        rels=SafeET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        for rel in rels:
            if rel.attrib.get("Type","").endswith("/worksheet"):
                target=rel.attrib.get("Target","")
                normalized=posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/"+target)
                if rel.attrib.get("TargetMode")=="External" or normalized not in sheets:
                    raise ValueError("Unsupported worksheet relationship.")
        merged_area=0
        for name in sheets:
            coords = []
            with archive.open(name) as stream:
                for _, node in SafeET.iterparse(stream, events=("end",)):
                    if time.monotonic()-started > 20:
                        raise ValueError("Workbook parsing exceeded its time limit.")
                    kind=node.tag.rsplit("}",1)[-1]
                    if kind=="mergeCell":
                        from openpyxl.utils.cell import range_boundaries
                        left,top,right,bottom=range_boundaries(node.attrib.get("ref",""))
                        merged_area+=(right-left+1)*(bottom-top+1)
                        if merged_area>MAX_CELLS:
                            raise ValueError("Workbook merged ranges exceed safe expansion limits.")
                    if kind != "c":
                        continue
                    physical += 1
                    if physical > MAX_PHYSICAL_CELLS:
                        raise ValueError("Workbook exceeds the physical cell traversal limit.")
                    if any(child.tag.rsplit("}",1)[-1] in ("v","f","is") and (child.text or len(child)) for child in node):
                        useful += 1
                        if useful > MAX_CELLS:
                            raise ValueError("Workbook exceeds 100,000 useful cells.")
                        coord = node.attrib.get("r","")
                        column, row = coordinate_from_string(coord)
                        col = column_index_from_string(column)
                        if not (1 <= row <= 1048576 and 1 <= col <= 16384):
                            raise ValueError("Invalid workbook cell coordinate.")
                        coords.append((row,col))
                    node.clear()
            coordinates[name] = coords
    # Preflight bounds actual serialized cells before openpyxl constructs objects.
    # Paired non-readonly maps permit sparse coordinate access without iterating
    # a rectangular max_row/max_column region or trusting dimension metadata.
    formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    cached = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    sections = []
    try:
        if len(formulas.worksheets) > MAX_SHEETS:
            raise ValueError("Workbook exceeds 20 sheets.")
        for sheet, values in zip(formulas.worksheets, cached.worksheets):
            # The bounded object maps contain physically present cells, not
            # dimension-generated blank cells. No rectangular worksheet iteration.
            cells = [(r,c,cell) for (r,c),cell in sheet._cells.items() if cell.value is not None]
            if len(cells)>MAX_CELLS:
                raise ValueError("Workbook exceeds useful cell limit.")
            cols = sorted({c for r,c,cell in cells})
            if len(cols)>MAX_COLUMNS:
                raise ValueError("A table exceeds 100 useful columns.")
            if not cells:
                continue
            by_row = defaultdict(dict)
            for r,c,cell in cells:
                by_row[r][c] = cell
            first = min(by_row)
            headers = []
            for c in cols:
                raw = by_row[first].get(c)
                headers.append(str(raw.value).strip() if raw is not None else f"Column {c}")
            if len(set(headers)) != len(headers):
                raise ValueError(f"Sheet {sheet.title} has duplicate column headings.")
            sections.append({"text":" | ".join(headers),"type":"table_header","section":sheet.title,
                             "sheet":sheet.title,"row":first,"headers":headers,"table":sheet.title})
            for r in sorted(by_row):
                if r == first:
                    continue
                row, formula_map = {}, {}
                for c,h in zip(cols,headers):
                    cell = by_row[r].get(c)
                    value = None if cell is None else cell.value
                    if cell is not None and cell.data_type == "f":
                        formula_map[h] = value
                        value = values.cell(r,c).value
                    row[h] = value.isoformat() if hasattr(value,"isoformat") else value
                sections.append({"text":" | ".join("" if row[h] is None else str(row[h]) for h in headers),
                    "type":"row","section":f"{sheet.title}, row {r}","sheet":sheet.title,"row":r,
                    "table":sheet.title,"headers":headers,"cells":row,"formulas":formula_map})
    finally:
        formulas.close(); cached.close()
    return sections

def tables_from_blocks(blocks):
    tables = {}
    for block in blocks:
        if block.get("type") not in ("row","table_header"):
            continue
        key = block.get("table","CSV")
        table = tables.setdefault(key, {"name":key,"headers":block.get("headers",[]),"rows":[]})
        if block["type"] == "row":
            table["rows"].append(block)
    return tables

def number(value):
    if isinstance(value, bool) or value is None or str(value).strip()=="":
        return None
    text = str(value).strip()
    if len(text)>300:
        raise ValueError("Numeric value exceeds safe precision limit.")
    exponent=re.search(r"[eE]([+-]?\d+)$",text)
    if exponent and abs(int(exponent[1]))>200:
        raise ValueError("Numeric exponent exceeds safe precision limit.")
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?",text):
        raise ValueError(f"Ambiguous/non-numeric value: {text[:60]}")
    try:
        result = Decimal(text)
    except InvalidOperation as error:
        raise ValueError("Invalid numeric value.") from error
    if not result.is_finite():
        raise ValueError("Non-finite numeric value.")
    return result

def analyze(blocks,spec):
    with localcontext() as context:
        context.prec=1000
        return _analyze(blocks,spec)

def _analyze(blocks, spec):
    if not isinstance(spec,dict) or set(spec)-{"table","group_by","filters","aggregates"}:
        raise ValueError("Invalid table operation.")
    if not all(isinstance(spec.get(k,[]),list) for k in ("group_by","filters","aggregates")):
        raise ValueError("Invalid table operation lists.")
    if any(not isinstance(v,dict) for k in ("filters","aggregates") for v in spec.get(k,[])):
        raise ValueError("Invalid filter or aggregate.")
    tables = tables_from_blocks(blocks)
    name = spec.get("table")
    if name is None and len(tables)==1:
        name = next(iter(tables))
    if name not in tables:
        raise ValueError("Choose a table/sheet: "+", ".join(tables))
    table = tables[name]
    columns = table["headers"]
    groups = spec.get("group_by",[])
    filters = spec.get("filters",[])
    aggregates = spec.get("aggregates",[])
    if len(groups)>3 or len(filters)>10 or not 1<=len(aggregates)<=8:
        raise ValueError("Table request exceeds operation limits.")
    def column(value):
        if value not in columns:
            raise ValueError("Unknown column; choose from: "+", ".join(columns))
    for key in groups:
        column(key)
    for f in filters:
        column(f.get("column"))
        if f.get("op") not in ("eq","ne","gt","gte","lt","lte","contains"):
            raise ValueError("Unsupported table filter.")
    for a in aggregates:
        if a.get("op") not in ("sum","average","min","max","count"):
            raise ValueError("Unsupported aggregate.")
        if a.get("column") is not None:
            column(a["column"])
        elif a.get("op")!="count":
            raise ValueError("An aggregate column is required.")
    selected = []
    formula_examples=[];formula_count=0
    missing = excluded = 0
    relevant = set(groups)|{f["column"] for f in filters}|{a["column"] for a in aggregates if a.get("column")}
    for row in table["rows"]:
        absent = [c for c in relevant if c in row.get("formulas",{}) and row["cells"].get(c) is None]
        if absent:
            raise ValueError(f"Formula has no cached value in {name}, row {row['row']}, columns {', '.join(absent)}. Recalculate and save it in your spreadsheet application.")
        for name_in_row in sorted(relevant):
            if name_in_row in row.get("formulas",{}):
                formula_count+=1
                if len(formula_examples)<200:
                    formula_examples.append({"sheet":name,"row":row["row"],"column":name_in_row,
                        "expression":row["formulas"][name_in_row],"cached_value":row["cells"].get(name_in_row)})
        keep = True
        for f in filters:
            value = row["cells"].get(f["column"])
            wanted = f.get("value")
            op = f["op"]
            if op in ("eq","ne","contains"):
                left,right=str("" if value is None else value).casefold(),str(wanted).casefold()
                passed = left==right if op=="eq" else left!=right if op=="ne" else right in left
            else:
                left,right=number(value),number(wanted)
                passed = left is not None and right is not None and {
                    "gt":lambda:left>right,"gte":lambda:left>=right,"lt":lambda:left<right,"lte":lambda:left<=right}[op]()
            keep &= passed
        if keep: selected.append(row)
    buckets=defaultdict(list)
    for row in selected:
        buckets[tuple(str("" if row["cells"].get(g) is None else row["cells"][g]) for g in groups)].append(row)
    if not groups and not buckets:
        buckets[()] = []
    results=[]
    if len(buckets)>200:
        raise ValueError("More than 200 result groups; narrow the filters.")
    for keys,rows in sorted(buckets.items()):
        output={"group":dict(zip(groups,keys)),"matched_rows":len(rows),"values":[]}
        for a in aggregates:
            op,c=a["op"],a.get("column")
            raw=[r["cells"].get(c) for r in rows] if c else [1]*len(rows)
            valid=[]
            for value in raw:
                if value is None or str(value).strip()=="":
                    missing+=1
                    continue
                if op=="count":
                    valid.append(1)
                else:
                    try:
                        numeric=number(value)
                        if numeric is not None:valid.append(numeric)
                        else:excluded+=1
                    except ValueError:
                        excluded+=1
            if op=="count":result=len(valid)
            elif not valid:result=None
            elif op=="sum":result=str(sum(valid,Decimal(0)))
            elif op=="average":result=str(sum(valid,Decimal(0))/len(valid))
            elif op=="min":result=str(min(valid))
            else:result=str(max(valid))
            output["values"].append({"operation":op,"column":c,"result":result})
        results.append(output)
    return {"table":name,"arguments":spec,"rows_examined":len(table["rows"]),
            "cells_examined":len(table["rows"])*len(relevant),
            "columns_examined":sorted(relevant),"formula_cells":formula_count,"formula_examples":formula_examples,
            "matched_rows":len(selected),"missing_values":missing,"excluded_values":excluded,
            "coverage":"complete","results":results}
