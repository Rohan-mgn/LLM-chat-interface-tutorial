"""Offline document-intelligence tests. No downloads or live model calls."""
import asyncio
from contextlib import closing
import io
import json
from pathlib import Path
import re
import sys
import unittest
from unittest.mock import patch, AsyncMock
from uuid import uuid4
import zipfile

STAGE=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(STAGE));sys.path.insert(0,str(STAGE/"tests"))
import check_rag as fixtures
from check_rag import FakeModel
from document_tools import *
from table_tools import workbook_sections, analyze
from documents import extract
from document_agent import rules
from summarizer import Summarizer
from model_provider import ModelProvider
from context_budget import check, estimate
from rag import Rag

class ExactTests(unittest.TestCase):
    def blocks(self,sections):return stable_blocks("fixture",sections)
    def test_unicode_whole_word_phrase_offsets(self):
        blocks=self.blocks([{"text":"Transformer transformers TRANSFORMER Straße neural\n network."}])
        self.assertEqual(count_word(blocks,"transformer")["count"],2)
        self.assertEqual(count_word(blocks,"strasse")["matches"][0]["text"],"Straße")
        self.assertEqual(count_phrase(blocks,"neural network")["count"],1)
        self.assertEqual(count_word(blocks,"transformer",True)["count"],0)
    def test_structural_boundaries_and_explicit_continuation(self):
        for boundary in ("page","row","sheet","section"):
            blocks=self.blocks([{"text":"neural",boundary:1},{"text":"network",boundary:2}])
            self.assertEqual(count_phrase(blocks,"neural network")["count"],0)
        blocks=self.blocks([{"text":"neural","logical_group":"p"},{"text":"network","logical_group":"p","continues_previous":True}])
        match=count_phrase(blocks,"neural network")
        self.assertEqual(match["count"],1);self.assertEqual(len(match["matches"][0]["locations"]),2)
    def test_pages_sections_extraction_and_statistics(self):
        b=self.blocks([{"text":"Heading","page":1,"section":"Intro","type":"heading","heading":"Heading"},
            {"text":"Write a@b.com https://example.org.","page":2,"section":"Contact"}])
        self.assertEqual(find_pages_containing(b,"write")["pages"],[2])
        self.assertEqual(get_page(b,2)[0]["section"],"Contact")
        self.assertEqual(get_section(b,"contact")[0]["page"],2)
        self.assertEqual(len(list_headings(b)),1)
        self.assertEqual(extract_emails(b)[0]["value"],"a@b.com")
        self.assertEqual(extract_urls(b)[0]["value"],"https://example.org")
        self.assertEqual(document_statistics(b)["pages"],2)
    def test_block_identifiers_stable_but_content_sensitive(self):
        a=self.blocks([{"text":"same","page":1}])
        self.assertEqual(a,self.blocks([{"text":"same","page":1}]))
        self.assertNotEqual(a[0]["id"],self.blocks([{"text":"changed","page":1}])[0]["id"])
    def test_router_acceptance_and_no_arithmetic_fallback(self):
        samples={"How many times does the word transformer occur?":"COUNT",
            "How many times is transformer mentioned?":"COUNT",
            "Which pages contain neural network?":"PAGE_SEARCH",
            "Summarize this 150-page document.":"SUMMARIZE_DOCUMENT",
            "Summarize section Introduction":"SUMMARIZE_SECTION",
            "Compare old_contract.pdf with new_contract.pdf.":"COMPARE_DOCUMENTS",
            "What is total revenue and average revenue by region?":"TABLE_ANALYSIS",
            "Why did the authors select transformer architecture?":"DOCUMENT_QA",
            "What about Q4?":"DOCUMENT_QA","Hello":"NORMAL_CHAT"}
        for text,route in samples.items():self.assertEqual(rules(text).route,route,text)
        self.assertEqual(rules("How many times does the word transformer occur?").term,"transformer")
    def test_budget_unicode_and_oversize(self):
        self.assertGreater(estimate("🙂"*100),estimate("a"*100))
        with self.assertRaisesRegex(ValueError,"budget"):check([{"role":"user","content":"🙂"*7000}])

class WorkbookTests(unittest.TestCase):
    def workbook(self,formula=False,styled=False):
        import openpyxl
        from openpyxl.styles import Font
        book=openpyxl.Workbook();sheet=book.active;sheet.title="Sales"
        sheet.append(["Region","Revenue"]);sheet.append(["East",10]);sheet.append(["East","=B2*2" if formula else 20]);sheet.append(["West",0])
        if styled:sheet["XFD1048576"].font=Font(bold=True)
        stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()
    def spec(self):return {"table":"Sales","group_by":["Region"],"filters":[],"aggregates":[{"op":"sum","column":"Revenue"},{"op":"average","column":"Revenue"}]}
    def test_xlsx_grouping_and_zero_not_missing(self):
        blocks=workbook_sections(self.workbook(styled=True))
        result=analyze(blocks,self.spec())
        self.assertEqual(result["rows_examined"],3)
        self.assertEqual(result["results"][0]["values"][0]["result"],"30")
        self.assertEqual(result["results"][1]["values"][0]["result"],"0")
    def test_formula_expression_and_absent_cache(self):
        blocks=workbook_sections(self.workbook(formula=True))
        row=next(b for b in blocks if b.get("row")==3)
        self.assertEqual(row["formulas"]["Revenue"],"=B2*2")
        self.assertIsNone(row["cells"]["Revenue"])
        with self.assertRaisesRegex(ValueError,"no cached value"):analyze(blocks,self.spec())
    def test_cached_formula_used_without_evaluation(self):
        source=zipfile.ZipFile(io.BytesIO(self.workbook(formula=True)))
        output=io.BytesIO()
        with source,zipfile.ZipFile(output,"w") as dest:
            for name in source.namelist():
                data=source.read(name)
                if name=="xl/worksheets/sheet1.xml":data=re.sub(rb"(<f>B2\*2</f>)<v(?:\s*/>|></v>)",rb"\1<v>999</v>",data)
                dest.writestr(name,data)
        result=analyze(workbook_sections(output.getvalue()),self.spec())
        self.assertEqual(result["results"][0]["values"][0]["result"],"1009")
    def test_useful_cell_and_merge_limits(self):
        with patch("table_tools.MAX_CELLS",4):
            with self.assertRaisesRegex(ValueError,"useful cells"):workbook_sections(self.workbook())
        import openpyxl
        book=openpyxl.Workbook();book.active.merge_cells("A1:ZZ1000")
        data=io.BytesIO();book.save(data);book.close()
        with self.assertRaisesRegex(ValueError,"merged ranges"):workbook_sections(data.getvalue())
    def test_invalid_columns_and_operators_are_not_code(self):
        blocks=workbook_sections(self.workbook())
        for spec in ({**self.spec(),"aggregates":[{"op":"eval","column":"Revenue"}]},
                     {**self.spec(),"aggregates":[{"op":"sum","column":"__import__('os')"}]}):
            with self.assertRaises(ValueError):analyze(blocks,spec)

class Integration(unittest.TestCase):
    setUp=fixtures.Checks.setUp;tearDown=fixtures.Checks.tearDown
    upload=fixtures.Checks.upload;send=fixtures.Checks.send;rows=fixtures.Checks.rows
    def test_exact_count_and_scoped_tool_evidence_without_model(self):
        file=self.upload(data=b"Transformer transformer transformers.")
        events=self.send("How many times does the word transformer occur?",use_files=True,file_ids=[file["id"]])
        self.assertEqual(events[-1]["status"],"completed",events)
        row=self.rows()[-1];self.assertIn("**2**",row["content"]);self.assertEqual(row["sources"],[])
        self.assertFalse(self.calls)
        evidence=row["tool_evidence"][0]["id"]
        response=self.client.get(f"/conversations/{self.cid}/tool-evidence/{evidence}")
        self.assertEqual(response.json()["exact_result"]["count"],2)
        self.assertEqual(response.json()["fingerprint"],file["sha256"])
        other=self.client.post("/conversations").json()["id"]
        self.assertEqual(self.client.get(f"/conversations/{other}/tool-evidence/{evidence}").status_code,404)
        self.client.delete(f"/conversations/{self.cid}/files/{file['id']}")
        self.assertEqual(self.client.get(f"/conversations/{self.cid}/tool-evidence/{evidence}").status_code,404)
    def test_use_files_false_bypasses_router_and_tools(self):
        file=self.upload()
        with patch.object(self.app.document_agent,"route",side_effect=AssertionError("must bypass")):
            events=self.send("How many times does falcon occur?",use_files=False,file_ids=[file["id"]])
        self.assertEqual(events[-1]["status"],"completed")
        self.assertNotIn("Blue Falcon",str(self.calls))
    def test_parsed_tools_survive_embedding_failure(self):
        with patch.object(FakeModel,"embed",side_effect=RuntimeError("embedding offline")):
            file=self.upload(data=b"transformer transformer",ready=False)
            asyncio.run(self.app.rag.index(self.cid,file["id"]))
        row=self.app.documents.row(self.cid,file["id"])
        self.assertEqual(row["parse_state"],"ready");self.assertEqual(row["state"],"failed")
        events=self.send("Count word transformer",use_files=True,file_ids=[file["id"]])
        self.assertEqual(events[-1]["status"],"completed",events);self.assertIn("**2**",self.rows()[-1]["content"])
    def test_csv_calculation_and_provenance(self):
        file=self.upload("sales.csv",b"region,revenue\nEast,0.1\nEast,0.2\nWest,\n","text/csv")
        spec={"table":"CSV","group_by":["region"],"filters":[],"aggregates":[{"op":"sum","column":"revenue"}]}
        with patch.object(self.app.provider,"structured",AsyncMock(return_value=spec)):
            events=self.send("What is total revenue by region?",use_files=True,file_ids=[file["id"]])
        self.assertEqual(events[-1]["status"],"completed",events)
        self.assertIn("0\\.3",self.rows()[-1]["content"])
        self.assertTrue(self.rows()[-1]["tool_evidence"])
    def test_fts_lexical_fallback_and_pipeline_scope(self):
        file=self.upload(data=b"Unique zebracircuit calibration settings.")
        async def run():
            client=self.app.provider.client()
            try:return await self.app.rag.retrieve(client,self.cid,[file["id"]],"zebracircuit")
            finally:await client.close()
        self.app.documents.fts_available=False
        with patch.object(FakeModel,"embed",side_effect=RuntimeError("offline")):
            hits,_=asyncio.run(run())
        self.assertTrue(hits);self.assertEqual(hits[0]["file_id"],file["id"])
        self.assertEqual(hits[0]["retrieval_backend"],"lexical")
        self.assertGreater(hits[0]["lexical_overlap"],0)
    def test_hybrid_fusion_reranker_fallback_and_diversity(self):
        rows=[{"id":"a","file_id":"1","text":"same same","dense_score":.8,"lexical_overlap":1},
            {"id":"b","file_id":"1","text":"same same","dense_score":.9,"lexical_overlap":1},
            {"id":"c","file_id":"2","text":"different relevant evidence","dense_score":.8,"lexical_overlap":1}]
        fused=Rag.fuse_results(rows,[rows[2],rows[0]])
        self.assertGreater(fused[0]["score"],1/61)
        diversified=Rag.diversify(fused,balanced=True)
        self.assertEqual({r["file_id"] for r in diversified},{"1","2"});self.assertEqual(len(diversified),2)
        with patch.object(self.app.provider,"structured",AsyncMock(side_effect=ValueError("bad JSON"))):
            self.assertEqual(asyncio.run(self.app.rag.rerank("question",fused)),fused[:12])
    def test_first_appearance_order_and_stable_reindex(self):
        file=self.upload(data=b"Falcon project first.\n\nFalcon project second.")
        with closing(self.app.get_connection()) as db:
            chunks=[dict(r) for r in db.execute("SELECT * FROM chunks WHERE file_id=? ORDER BY chunk_index",(file["id"],))]
        self.send("hello")
        mid=self.rows()[-1]["id"]
        text="["+chunks[1]["id"]+"] then ["+chunks[0]["id"]+"]"
        self.app.save_generation(mid,text,"completed")
        self.app.rag.persist(mid,text,{"sources":chunks,"debug":{"context":"private duplicate"}})
        self.assertEqual([s["id"] for s in self.rows()[-1]["sources"]],[chunks[1]["id"],chunks[0]["id"]])
        asyncio.run(self.app.rag.index(self.cid,file["id"]))
        self.assertEqual([s["id"] for s in self.rows()[-1]["sources"]],[chunks[1]["id"],chunks[0]["id"]])
        with closing(self.app.get_connection()) as db:
            self.assertNotIn("context",json.loads(db.execute("SELECT details FROM rag_runs WHERE message_id=?",(mid,)).fetchone()[0]))
    def test_migration_preserves_history_and_citations_backup_repeatable(self):
        file=self.upload();self.send(use_files=True,file_ids=[file["id"]])
        before=self.rows();source=before[-1]["sources"][0]["id"]
        with closing(self.app.get_connection()) as db,db:
            db.execute("DELETE FROM document_schema")
            db.execute("UPDATE files SET parser_config=NULL WHERE id=?",(file["id"],))
        self.app.init_db();self.app.init_db()
        self.assertTrue((self.root/"chat.pre-document-intelligence.db").is_file())
        self.assertEqual(self.rows()[-1]["content"],before[-1]["content"])
        self.assertEqual(self.app.documents.row(self.cid,file["id"])["state"],"reindex_required")
        self.assertEqual(self.client.get(f"/conversations/{self.cid}/sources/{source}").status_code,200)
        with closing(self.app.get_connection()) as db:self.assertFalse(db.execute("PRAGMA foreign_key_check").fetchall())
    def test_two_conversations_and_same_conversation_conflict(self):
        other=self.client.post("/conversations").json()["id"]
        original=self.app.document_agent.prepare
        async def run():
            entered=asyncio.Event();release=asyncio.Event()
            async def prepare(job,messages):
                if job["conversation_id"]==self.cid:
                    entered.set();await release.wait()
                return await original(job,messages)
            with patch.object(self.app.document_agent,"prepare",side_effect=prepare):
                first=await self.app.chat(self.app.ChatRequest(conversation_id=self.cid,message="Long operation",generation_id=uuid4()))
                await entered.wait()
                with self.assertRaises(Exception) as error:
                    await self.app.chat(self.app.ChatRequest(conversation_id=self.cid,message="Conflict",generation_id=uuid4()))
                self.assertEqual(error.exception.status_code,409)
                second=await self.app.chat(self.app.ChatRequest(conversation_id=other,message="Independent",generation_id=uuid4()))
                self.assertEqual(await asyncio.wait_for(second.job["task"],2),"completed")
                self.assertFalse(first.job["task"].done())
                release.set();await first.job["task"]
            self.assertFalse(self.app.conversation_lock.locked())
        asyncio.run(run())
    def test_vision_work_does_not_block_native_indexing(self):
        from pypdf import PdfWriter
        writer=PdfWriter();writer.add_blank_page(width=200,height=200)
        data=io.BytesIO();writer.write(data)
        scanned=self.upload("scan.pdf",data.getvalue(),"application/pdf",ready=False)
        ordinary=self.upload("ordinary.txt",b"Native text content.",ready=False)
        # Background upload work finishes first; the controlled reindex uses two
        # tasks with a deliberately stalled vision call and no global index gate.
        import time
        time.sleep(.1)
        async def run():
            entered=asyncio.Event();release=asyncio.Event()
            async def vision(path,sections,*args):
                entered.set();await release.wait()
                return sections,"Vision unavailable"
            with patch("rag.vision_sections",side_effect=vision):
                slow=asyncio.create_task(self.app.rag.index(self.cid,scanned["id"]))
                await asyncio.wait_for(entered.wait(),2)
                await asyncio.wait_for(self.app.rag.index(self.cid,ordinary["id"]),2)
                self.assertEqual(self.app.documents.row(self.cid,ordinary["id"])["state"],"ready")
                self.assertFalse(slow.done());release.set();await slow
        asyncio.run(run())
    def test_summary_full_coverage_section_cache_and_limits(self):
        file=self.upload("summary.md",b"# First\n\nAlpha text.\n\n# Last\n\nOmega final fact.","text/markdown")
        calls=[]
        class Client:
            async def chat(inner,**kwargs):
                calls.append(kwargs)
                ids=re.findall(r'SOURCE_[a-f0-9]+',str(kwargs["messages"]))
                return {"message":{"content":"Faithful summary "+("["+ids[0]+"]" if ids else "")}}
            async def close(inner):pass
        async def run():
            with patch.object(self.app.provider,"client",return_value=Client()):
                docs,sources,info=await self.app.document_agent.summarizer.summarize(self.cid,[file],lambda _:None)
                self.assertIn("Omega final fact",str(calls));self.assertEqual(info["calls"],1)
                await self.app.document_agent.summarizer.summarize(self.cid,[file],lambda _:None)
                self.assertEqual(len(calls),1)
                calls.clear()
                await self.app.document_agent.summarizer.summarize(self.cid,[file],lambda _:None,section="Last")
                self.assertIn("Omega",str(calls));self.assertNotIn("Alpha",str(calls))
                with patch("summarizer.MAX_CALLS",0):
                    with self.assertRaisesRegex(ValueError,"64-call"):
                        await self.app.document_agent.summarizer.summarize(self.cid,[file],lambda _:None,section="First")
        asyncio.run(run())
    def test_second_retrieval_round_cannot_loop(self):
        file=self.upload()
        async def structure(messages,schema,**kwargs):
            if "query" in schema.get("properties",{}):return {"query":"project falcon missing reasoning"}
            raise ValueError("fallback rerank")
        with patch.object(self.app.provider,"structured",side_effect=structure),patch.object(self.app.rag,"retrieve",wraps=self.app.rag.retrieve) as retrieve:
            events=self.send("Why did the project choose Falcon?",use_files=True,file_ids=[file["id"]])
            self.assertEqual(events[-1]["status"],"completed",events);self.assertEqual(retrieve.call_count,2)

    def test_vision_cache_native_preservation_and_no_vision_dependency(self):
        from document_parsers import vision_sections
        cache={};requests=[]
        class Client:
            async def show(self,model):return {"capabilities":["vision"],"model_info":{"version":"fixture"}}
            async def close(self):pass
        class Provider:
            vision_model="fixture-vision"
            async def vision_available(self):return True,""
            def client(self,*args):return Client()
            async def structured(self,*args,**kwargs):
                requests.append(kwargs)
                return {"transcription":"Scanned actual words","description":"A page","uncertainty":""}
        async def run():
            sections=[{"text":"Native text stays","page":1,"type":"page_text"},
                {"text":"","page":2,"type":"unreadable"}]
            with patch("document_parsers.render_page",return_value=b"PNG-fixture"):
                result,warning=await vision_sections(Path("unused.pdf"),sections,Provider(),cache.get,lambda k,v:cache.update({k:v}))
                self.assertEqual(result[0]["text"],"Native text stays")
                self.assertEqual(requests[0]["images"],[b"PNG-fixture"])
                self.assertEqual(result[1]["origin"],"vision")
                self.assertEqual(result[1]["page"],2)
                await vision_sections(Path("unused.pdf"),sections,Provider(),cache.get,lambda k,v:cache.update({k:v}))
                self.assertEqual(len(requests),1)
        asyncio.run(run())

    def test_hierarchical_summary_covers_end_and_cancellation(self):
        file=self.upload("large.txt",b"placeholder")
        sections=[{"text":f"Unique section {i}. "+("Useful content. "*65),"section":f"Section {i}"} for i in range(55)]
        sections[-1]["text"]+=" CRITICAL_FINAL_FACT"
        self.app.documents.store_extracted(self.cid,file["id"],sections)
        calls=[]
        class Client:
            async def chat(inner,**kwargs):
                calls.append(kwargs)
                ids=re.findall(r"SOURCE_[a-f0-9]+",str(kwargs["messages"]))
                return {"message":{"content":"Summary "+("["+ids[-1]+"]" if ids else "")}}
            async def close(inner):pass
        async def run():
            with patch.object(self.app.provider,"client",return_value=Client()):
                docs,sources,info=await self.app.document_agent.summarizer.summarize(self.cid,[file],lambda _:None)
            self.assertGreater(info["calls"],1);self.assertLess(info["calls"],15)
            self.assertEqual(len(sources),55);self.assertIn("CRITICAL_FINAL_FACT",str(calls))
            self.assertEqual(docs[0]["blocks_read"],55)
            self.assertEqual(len(sources),len(set(s["id"] for s in sources)))
            entered=asyncio.Event();closed=[]
            class Waiting:
                async def chat(inner,**kwargs):entered.set();await asyncio.sleep(60)
                async def close(inner):closed.append(True)
            with patch.object(self.app.provider,"client",return_value=Waiting()):
                task=asyncio.create_task(self.app.document_agent.summarizer.summarize(self.cid,[file],lambda _:None,section="Section 54"))
                await asyncio.wait_for(entered.wait(),1);task.cancel()
                with self.assertRaises(asyncio.CancelledError):await task
                self.assertTrue(closed)
        asyncio.run(run())

    def test_summary_comparison_uses_every_selected_file(self):
        a=self.upload("a.txt",b"Contract old: term 30 days.")
        b=self.upload("b.txt",b"Contract new: term 60 days.")
        seen=[]
        class Client:
            async def chat(inner,**kwargs):
                seen.append(str(kwargs["messages"]))
                ids=re.findall(r"SOURCE_[a-f0-9]+",str(kwargs["messages"]))
                return {"message":{"content":"Terms differ ["+ids[0]+"]"}}
            async def close(inner):pass
        async def run():
            with patch.object(self.app.provider,"client",return_value=Client()):
                docs,sources,_=await self.app.document_agent.summarizer.summarize(self.cid,[a,b],lambda _:None)
            self.assertEqual({s["file_id"] for s in sources},{a["id"],b["id"]})
            self.assertIn("30 days",str(seen));self.assertIn("60 days",str(seen))
        asyncio.run(run())

    def test_history_compression_is_branch_keyed(self):
        messages=[{"role":"user","content":("long past context "*300)+str(i)} for i in range(3)]+[{"role":"user","content":"Recent question"}]
        calls=[]
        class Client:
            async def chat(inner,**kwargs):
                calls.append(kwargs);return {"message":{"content":"Unverified past user facts"}}
            async def close(inner):pass
        async def run():
            with patch.object(self.app.provider,"client",return_value=Client()):
                result=await self.app.document_agent.history(self.cid,messages)
                self.assertEqual(result[-1],messages[-1]);self.assertIn("Unverified",result[0]["content"])
                count=len(calls)
                await self.app.document_agent.history(self.cid,messages)
                self.assertEqual(len(calls),count)
                edited=[{**messages[0],"content":"Different branch"},*messages[1:]]
                await self.app.document_agent.history(self.cid,edited)
                self.assertGreater(len(calls),count)
        asyncio.run(run())

class ProviderTests(unittest.TestCase):
    def test_default_one_configurable_two_and_cancelled_waiter(self):
        async def run(limit):
            with patch.dict("os.environ",{"TEXT_MODEL_CONCURRENCY":str(limit)}):
                provider=ModelProvider()
            self.assertEqual(provider.text_concurrency,limit)
            active=0;peak=0;release=asyncio.Event()
            async def work():
                nonlocal active,peak
                async with provider.slot(provider.text_slots):
                    active+=1;peak=max(peak,active)
                    try:await release.wait()
                    finally:active-=1
            tasks=[asyncio.create_task(work()) for _ in range(3)]
            await asyncio.sleep(.05)
            self.assertEqual(peak,limit)
            tasks[-1].cancel()
            await asyncio.gather(tasks[-1],return_exceptions=True)
            release.set();await asyncio.gather(*tasks[:-1])
            async with provider.slot(provider.text_slots):pass
        asyncio.run(run(1));asyncio.run(run(2))

if __name__=="__main__":unittest.main(verbosity=2)
